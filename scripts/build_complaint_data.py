#!/usr/bin/env python3
"""将投诉登记转换为不含客户信息的看板轻量数据。"""
from __future__ import annotations

import argparse
import html
import json
import re
import subprocess
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "data" / "complaint-data.js"

TYPE_NAMES = ["产品质量与性能", "价格与促销", "其他/待核实", "包装问题", "施工与使用指导", "服务沟通", "物流与发货", "退换与赔付诉求"]
TYPE_RULES = [
    ("服务沟通", re.compile(r"客服|态度|回复|沟通|联系不上|无人处理|推诿|服务")),
    ("物流与发货", re.compile(r"物流|快递|发货|少发|漏发|错发|丢件|运输|送货")),
    ("包装问题", re.compile(r"包装|漏液|渗漏|破桶|爆桶|漏桶|桶破|破损|撒漏")),
    ("价格与促销", re.compile(r"价格|价保|降价|优惠|赠品|活动|差价")),
    ("产品质量与性能", re.compile(r"质量|发霉|变质|不干|不固化|脱落|起皮|开裂|裂纹|漏水|渗水|不防水|色差|掉色|生锈|腐蚀|异味|堵塞|鼓包|起泡|失效|瑕疵")),
    ("施工与使用指导", re.compile(r"施工|使用|操作|涂刷|基层|兑水|加水|配比|比例|搅拌|养护|固砂|咨询|怎么用|用法")),
    ("退换与赔付诉求", re.compile(r"退货|退款|退换|换货|赔偿|赔付|补偿|补发")),
]


def plain(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def cell_text(cell) -> str:
    if not isinstance(cell, dict):
        return plain(cell)
    values = cell.get("data") or []
    return " ".join(plain(item.get("text") or item.get("name")) for item in values if isinstance(item, dict)).strip()


def classify(detail: str, product: str) -> str:
    joined = f"{detail} {product}"
    for label, pattern in TYPE_RULES:
        if pattern.search(joined):
            return label
    return "其他/待核实"


def clean_product(value: str) -> str:
    value = re.sub(r"\s+", " ", plain(value))
    if not value:
        return "未填写产品"
    if len(value) > 60 or "工单" in value or re.fullmatch(r"\d{8,}", value):
        return "未规范产品"
    return value


def parse_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = plain(value)
    match = re.search(r"(20\d{2})[/-](\d{1,2})[/-](\d{1,2})", text)
    if not match:
        return ""
    return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"


def clipboard_records():
    copied = subprocess.run(["pbpaste", "-Prefer", "html"], capture_output=True, text=True, check=True).stdout
    match = re.search(r"data-byte-bitable='([^']+)'", copied)
    if not match:
        raise RuntimeError("剪贴板中没有飞书多维表格数据；请先全选投诉记录并复制")
    payload = json.loads(html.unescape(match.group(1)))
    headers = [item.get("name", "") for item in payload.get("header", [])]
    rows = []
    for source in payload.get("body", []):
        record = {headers[i]: cell_text(source[i]) for i in range(min(len(headers), len(source)))}
        rows.append(record)
    return rows


def workbook_records(path: Path, sheet: str | None):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook[workbook.sheetnames[0]]
    headers = [plain(cell.value) for cell in worksheet[1]]
    return [{headers[i]: value for i, value in enumerate(values) if i < len(headers)} for values in worksheet.iter_rows(min_row=2, values_only=True)]


def platform_for(shop: str) -> str:
    if "天猫" in shop:
        return "天猫"
    if "京东" in shop:
        return "京东"
    if "拼多多" in shop:
        return "拼多多"
    if any(key in shop for key in ("抖音", "快手", "视频号", "小红书")):
        return "内容电商"
    return "其他渠道" if shop else ""


def week_key(value: date) -> str:
    year, week, _ = value.isocalendar()
    return f"{year}-W{week:02d}"


def week_start(key: str) -> date:
    year, week = key.split("-W")
    return date.fromisocalendar(int(year), int(week), 1)


def month_range(start: str, end: str):
    year, month = map(int, start.split("-"))
    end_year, end_month = map(int, end.split("-"))
    while (year, month) <= (end_year, end_month):
        yield f"{year:04d}-{month:02d}"
        month += 1
        if month == 13:
            year, month = year + 1, 1


def week_range(start: str, end: str):
    current, finish = week_start(start), week_start(end)
    while current <= finish:
        yield week_key(current)
        current += timedelta(days=7)


def build(records):
    safe_rows = []
    ticket_ids = []
    for record in records:
        ticket = plain(record.get("工单编号"))
        complaint_date = parse_date(record.get("投诉日期"))
        product = clean_product(record.get("产品"))
        detail = plain(record.get("投诉详情"))
        status = plain(record.get("工单状态")) or "未同步"
        result = 1 if plain(record.get("处理结果")) else 0
        shop = plain(record.get("店铺"))
        if ticket:
            ticket_ids.append(ticket)
        safe_rows.append({"ticket": ticket, "date": complaint_date, "type": classify(detail, product), "product": product,
                          "status": status, "result": result, "platform": platform_for(shop), "shop": shop})

    source_rows = len(safe_rows)
    seen_tickets = set()
    deduplicated = []
    for row in safe_rows:
        ticket = row["ticket"]
        if ticket and ticket in seen_tickets:
            continue
        if ticket:
            seen_tickets.add(ticket)
        deduplicated.append(row)
    safe_rows = deduplicated

    dates = sorted(row["date"] for row in safe_rows if row["date"])
    if not dates:
        raise RuntimeError("投诉日期全部为空，停止覆盖现有数据")
    parsed_rows = [(datetime.strptime(row["date"], "%Y-%m-%d").date(), row) for row in safe_rows if row["date"]]
    today = datetime.now().astimezone().date()
    latest_complete_end = today - timedelta(days=today.weekday() + 1)
    latest_complete_week = week_key(latest_complete_end)
    prior_complete_week = week_key(latest_complete_end - timedelta(days=7))
    current_partial_week = week_key(today)

    def period_rows(grain: str):
        grouped = defaultdict(list)
        for parsed_date, row in parsed_rows:
            grouped[parsed_date.strftime("%Y-%m") if grain == "month" else week_key(parsed_date)].append(row)
        keys = (month_range(dates[0][:7], today.strftime("%Y-%m")) if grain == "month"
        else week_range(week_key(datetime.strptime(dates[0], "%Y-%m-%d").date()), current_partial_week))
        output = []
        for key in keys:
            rows = grouped.get(key, [])
            type_counts = Counter(row["type"] for row in rows)
            output.append([key, len(rows), *[type_counts[name] for name in TYPE_NAMES],
                           sum(row["result"] for row in rows), sum(row["status"] != "未同步" for row in rows)])
        return output

    recent_weeks = list(week_range(week_key(latest_complete_end - timedelta(days=21)), latest_complete_week))
    product_by_week = defaultdict(Counter)
    channel_by_week = defaultdict(Counter)
    for parsed_date, row in parsed_rows:
        key = week_key(parsed_date)
        if key not in recent_weeks:
            continue
        if row["product"] not in ("未填写产品", "未规范产品"):
            product_by_week[key][row["product"]] += 1
        if row["platform"]:
            channel_by_week[key][row["platform"]] += 1

    def recent_rows(index):
        names = set().union(*(index[key].keys() for key in recent_weeks))
        rows = [[name, sum(index[key][name] for key in recent_weeks),
                 *[index[key][name] for key in recent_weeks]] for name in names]
        return sorted(rows, key=lambda row: (-row[1], row[0]))

    store_exported = sum(bool(row["shop"]) for row in safe_rows)
    return {
        "meta": {
            "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
            "sourceMin": dates[0] if dates else "",
            "sourceMax": dates[-1] if dates else "",
            "sourceRows": source_rows,
            "recordCount": len(safe_rows),
            "uniqueTickets": len(set(ticket_ids)),
            "duplicateTickets": len(ticket_ids) - len(set(ticket_ids)),
            "missingTicket": source_rows - len(ticket_ids),
            "missingDate": sum(not row["date"] for row in safe_rows),
            "missingProduct": sum(row["product"] == "未填写产品" for row in safe_rows),
            "unstandardizedProduct": sum(row["product"] == "未规范产品" for row in safe_rows),
            "missingResult": sum(not row["result"] for row in safe_rows),
            "missingStatus": sum(row["status"] == "未同步" for row in safe_rows),
            "storeExported": store_exported,
            "latestCompleteWeek": latest_complete_week,
            "priorCompleteWeek": prior_complete_week,
            "currentPartialWeek": current_partial_week,
            "storeFieldNote": "店铺为关联字段，当前账号仅可见标签，导出为空；平台/店铺/渠道暂不可完整统计" if not store_exported else "店铺关联字段可用",
            "classificationVersion": "投诉文本规则初分-v1.0-20260910",
        },
        "types": TYPE_NAMES,
        "months": period_rows("month"),
        "weeks": period_rows("week"),
        "recentProductPeriod": f"{recent_weeks[0]} 至 {recent_weeks[-1]}",
        "recentProductWeeks": recent_weeks,
        "recentProductWeekly": recent_rows(product_by_week),
        "recentChannelWeeks": recent_weeks,
        "recentChannelWeekly": recent_rows(channel_by_week),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    records = workbook_records(args.xlsx, args.sheet) if args.xlsx else clipboard_records()
    data = build(records)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text("window.COMPLAINT_DATA=" + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    print(json.dumps(data["meta"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
