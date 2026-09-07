#!/usr/bin/env python3
"""将售后登记与销售周报转换为客服看板所需的脱敏聚合数据。"""
from __future__ import annotations

import json
import math
import os
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
AFTERSALES = Path(os.environ.get("AFTERSALES_XLSX", "/Users/koralgrey/Downloads/YH售后问题&打款登记.xlsx"))
SALES = Path(os.environ.get("SALES_XLSX", "/Users/koralgrey/Library/Containers/com.tencent.xinWeChat/Data/Documents/xwechat_files/chenjiahui977450_57c0/msg/file/2026-08/周报(2).xlsx"))
OUTPUT_DIR = Path(os.environ.get("OUTPUT_DIR", str(ROOT / "data")))
OUTPUT = OUTPUT_DIR / "after-sales-data.js"
PRODUCT_OUTPUT = OUTPUT_DIR / "after-sales-product-sales.js"
SHOP_SALES_OUTPUT = OUTPUT_DIR / "after-sales-shop-sales.js"


def text(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def code_text(v) -> str:
    s = text(v)
    return re.sub(r"\.0$", "", s)


BRAND_ONLY_PRODUCT_NAMES = {"飞鱼", "东方雨虹", "雨虹"}


def product_display_name(row: dict, code: str) -> str:
    """优先使用运营产品简称，但品牌名不能作为商品名。"""
    operating_name = text(row.get("运营产品简称"))
    if operating_name and operating_name not in BRAND_ONLY_PRODUCT_NAMES:
        return operating_name
    return text(row.get("产品简称")) or text(row.get("产品描述")) or code


def number(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(float(v)) else 0.0
    s = re.sub(r"[,\s¥￥]", "", str(v))
    try:
        n = float(s)
        return n if math.isfinite(n) else 0.0
    except ValueError:
        return 0.0


def parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if not v:
        return None
    s = str(v).strip().replace("/", "-").replace(".", "-")
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def period_keys(d: date):
    iso = d.isocalendar()
    return f"{d.year}-{d.month:02d}", f"{iso.year}-W{iso.week:02d}"


def aftersales_platform(store: str):
    if "天猫" in store:
        return "天猫平台"
    if "京东" in store:
        return "京东平台"
    if "拼多多" in store:
        return "拼多多"
    if any(x in store for x in ("抖音", "快手", "视频号", "小红书", "微盟")):
        return "兴趣电商"
    return "其他电商"


SHOP_ALIASES = {
    # 天猫：售后表与销售周报的店铺别名
    ("天猫平台", "【天猫】东方雨虹官方旗舰店"): "东方雨虹天猫主店",
    ("天猫平台", "雨虹防水官方旗舰店"): "东方雨虹天猫主店",
    ("天猫平台", "【天猫】雨虹飞鱼旗舰店"): "雨虹飞鱼天猫店",
    ("天猫平台", "雨虹飞鱼天猫店"): "雨虹飞鱼天猫店",
    # 京东
    ("京东平台", "【京东自营】东方雨虹京东自营官方旗舰店"): "东方雨虹京东自营旗舰店",
    ("京东平台", "东方雨虹京东自营旗舰店"): "东方雨虹京东自营旗舰店",
    ("京东平台", "【京东】雨虹飞鱼旗舰店"): "雨虹飞鱼京东pop店",
    ("京东平台", "雨虹飞鱼京东pop店"): "雨虹飞鱼京东pop店",
    ("京东平台", "【京东】alpina阿尔贝娜旗舰店"): "alpina阿尔贝娜旗舰6068",
    ("京东平台", "alpina阿尔贝娜旗舰6068"): "alpina阿尔贝娜旗舰6068",
    # 拼多多
    ("拼多多", "【拼多多】东方雨虹旗舰店"): "东方雨虹旗舰店拼多多",
    ("拼多多", "东方雨虹旗舰店拼多多"): "东方雨虹旗舰店拼多多",
    ("拼多多", "【拼多多】雨虹飞鱼官方旗舰店"): "雨虹飞鱼拼多多店",
    ("拼多多", "雨虹飞鱼拼多多店"): "雨虹飞鱼拼多多店",
    # 兴趣电商：保留渠道区分，避免抖音/快手同名店误合并
    ("兴趣电商", "【抖音】东方雨虹官方旗舰店"): "抖音东方雨虹旗舰店",
    ("兴趣电商", "东方雨虹官方旗舰店1933"): "抖音东方雨虹旗舰店",
    ("兴趣电商", "【快手】东方雨虹官方旗舰店"): "快手东方雨虹旗舰店",
    ("兴趣电商", "雨虹旗舰店"): "快手东方雨虹旗舰店",
    ("兴趣电商", "【抖音】雨虹飞鱼官方旗舰店"): "抖音雨虹飞鱼旗舰店",
    ("兴趣电商", "雨虹飞鱼建材旗舰店抖音"): "抖音雨虹飞鱼旗舰店",
    ("兴趣电商", "【视频号】东方雨虹家用建材商城"): "视频号东方雨虹家用建材商城",
    ("兴趣电商", "视频号东方雨虹家用建材商城"): "视频号东方雨虹家用建材商城",
}


def shop_key(store: str, platform: str):
    raw = re.sub(r"[\s（）()_-]", "", text(store)).lower()
    alias = SHOP_ALIASES.get((platform, raw))
    if alias:
        return re.sub(r"[\s（）()_-]", "", alias.lower())
    s = re.sub(r"^【[^】]+】", "", raw)
    return s


RULES = [
    ("运输破损", ("运输破损", "物流破损", "快递破损", "暴力运输", "运输导致")),
    ("包装渗漏", ("渗漏", "泄漏", "漏液", "漏料", "漏胶", "漏粉", "漏水", "胀包")),
    ("损坏/破瓶", ("破瓶", "破损", "破裂", "碎了", "压瘪", "破桶", "破罐")),
    ("少发/漏发", ("少发", "漏发", "缺货", "少一", "没发", "未发")),
    ("错发/混发", ("错发", "混发", "发错", "错货", "不是订购")),
    ("配件问题", ("缺配件", "少配件", "配件坏", "配件不齐", "胶嘴", "喷头", "工具缺")),
    ("包装异常", ("包装", "外包", "内包", "标签", "瓶盖", "桶盖", "封口")),
    ("不干/不固化", ("不干", "不固化", "未固化", "固化慢", "不硬", "软胶")),
    ("爆胶/爆桶", ("爆胶", "爆桶", "爆罐", "喷胶", "胀桶")),
    ("开裂", ("开裂", "裂缝", "龟裂", "裂开")),
    ("脱落/掉皮", ("脱落", "掉皮", "起皮", "剥落", "掉块")),
    ("发霉/变质", ("发霉", "霉变", "变质", "结块", "变硬", "过期")),
    ("划痕/磕碰", ("划痕", "划伤", "磕碰", "磕伤", "凹陷", "掉漆")),
    ("变形", ("变形", "弯曲", "翘边", "不平")),
    ("生锈", ("生锈", "锈蚀", "长锈")),
    ("色差/颜色异常", ("色差", "颜色不对", "变色", "褪色", "色号", "颜色差")),
    ("异味", ("异味", "臭味", "味道大", "刺鼻")),
    ("防水效果异常", ("不防水", "防水失效", "漏水", "渗水", "返潮")),
    ("粘接性异常", ("粘不住", "不粘", "粘性", "附着力", "粘合")),
    ("效果不及预期", ("效果不好", "无效", "没效果", "不管用", "不如预期")),
    ("施工/使用问题", ("施工", "不会用", "使用方法", "操作", "打不出", "堵塞", "喷不出")),
    ("规格/数量不符", ("规格不符", "数量不符", "型号不对", "尺寸不对", "克重")),
    ("售后服务问题", ("客服", "服务态度", "回复慢", "不回复", "处理慢")),
    ("价格/活动问题", ("价差", "价格", "活动", "优惠券", "降价", "买贵")),
]


def classify(row: dict):
    joined = " ".join(text(row.get(k)) for k in ("问题大类", "问题描述", "问题描述2", "原因核实结果", "打款类型", "打款描述", "处理方式", "AI三级分类", "AI二级分类"))
    for label, keys in RULES:
        if any(k in joined for k in keys):
            return label, "高"
    for k in ("AI三级分类", "AI二级分类", "问题大类"):
        v = text(row.get(k))
        if v and v not in ("-", "其他", "产品问题"):
            return v[:30], "中"
    return "其他/待核实", "低"


def row_dict(headers, values):
    return {text(h): values[i] if i < len(values) else None for i, h in enumerate(headers) if h is not None}


def add(bucket, key, *, count=0.0, qty=0.0, amount=0.0, paid_amount=0.0, paid_count=0.0, max_single=0.0):
    r = bucket[key]
    if not r:
        r.update(dict(zip(key[0::2], key[1::2])))
    if count:
        r["count"] = r.get("count", 0) + count
    if qty:
        r["qty"] = r.get("qty", 0) + qty
    if amount:
        r["amount"] = r.get("amount", 0) + amount
    if paid_amount:
        r["paidAmount"] = r.get("paidAmount", 0) + paid_amount
    if paid_count:
        r["paidCount"] = r.get("paidCount", 0) + paid_count
    if max_single:
        r["maxSingle"] = max(r.get("maxSingle", 0), max_single)


def compact(bucket):
    rows = list(bucket.values())
    for r in rows:
        for k in ("count", "qty", "amount", "paidAmount", "paidCount", "maxSingle"):
            if k in r:
                r[k] = round(r[k], 6 if k == "amount" else 4)
    return rows


def load_product_maps(sales_wb, after_wb):
    ws = sales_wb["商品分类表"]
    headers = [text(v) for v in next(ws.iter_rows(values_only=True))]
    master = {}
    barcode = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        r = row_dict(headers, values)
        code = code_text(r.get("新编码"))
        if not code:
            continue
        info = {
            "code": code,
            "cat1": text(r.get("运营一级大类")) or "未分类",
            "cat2": text(r.get("运营二级品类")) or "未分类",
            "product": product_display_name(r, code),
            "spec": text(r.get("产品规格")),
            "barcode": code_text(r.get("69码")),
        }
        master[code] = info
        if info["barcode"]:
            barcode[info["barcode"]] = info

    ws = after_wb["产品源-优先引用"]
    headers = [text(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    name_to_code = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        r = row_dict(headers, values)
        name = text(r.get("产品名称"))
        code = code_text(r.get("产品编码（新）"))
        if name and code:
            name_to_code[name] = code
    return master, barcode, name_to_code


def main():
    after_wb = load_workbook(AFTERSALES, read_only=False, data_only=True)
    sales_wb = load_workbook(SALES, read_only=True, data_only=True)
    master, barcode_map, name_to_code = load_product_maps(sales_wb, after_wb)

    sales_m, sales_w = defaultdict(dict), defaultdict(dict)
    product_sales_m, product_sales_w = defaultdict(dict), defaultdict(dict)
    shop_sales_m, shop_sales_w = defaultdict(dict), defaultdict(dict)
    sales_max_date = None
    excluded_dealer_qty = excluded_dealer_amount = 0.0

    def consume_sales(sheet_name, legacy=False):
        nonlocal sales_max_date, excluded_dealer_qty, excluded_dealer_amount
        ws = sales_wb[sheet_name]
        it = ws.iter_rows(values_only=True)
        headers = [text(v) for v in next(it)]
        blank_run = 0
        for values in it:
            r = row_dict(headers, values)
            year = int(number(r.get("年份" if legacy else "年")))
            month = int(number(r.get("月")))
            week = int(number(r.get("销售周" if legacy else "周")))
            qty = number(r.get("数量"))
            amount = number(r.get("金额" if legacy else "实际金额"))
            code = code_text(r.get("雨虹编码" if legacy else "商品编码"))
            if not year and not month and not week and not code and not qty:
                blank_run += 1
                if blank_run >= 200:
                    break
                continue
            blank_run = 0
            if year < 2025 or month not in range(1, 13) or week not in range(1, 54) or qty <= 0:
                continue
            platform = text(r.get("平台")) or "其他电商"
            if platform == "经销商":
                excluded_dealer_qty += qty
                excluded_dealer_amount += max(0, amount)
                continue
            shop = text(r.get("店铺" if legacy else "店铺全称")) or text(r.get("店铺")) or "未填写"
            skey = shop_key(shop, platform)
            info = master.get(code, {})
            cat1 = text(r.get("运营一级大类")) or info.get("cat1") or "未分类"
            cat2 = text(r.get("运营二级品类")) or info.get("cat2") or "未分类"
            product = text(r.get("运营产品简称")) or info.get("product") or text(r.get("商品名称")) or code or "未匹配产品"
            m, w = f"{year}-{month:02d}", f"{year}-W{week:02d}"
            common = ("cat1", cat1, "cat2", cat2)
            add(sales_m, ("period", m) + common, qty=qty, amount=max(0, amount))
            add(sales_w, ("period", w) + common, qty=qty, amount=max(0, amount))
            shop_common = ("platform", platform, "shop", shop, "shopKey", skey) + common
            add(shop_sales_m, ("period", m) + shop_common, qty=qty, amount=max(0, amount))
            add(shop_sales_w, ("period", w) + shop_common, qty=qty, amount=max(0, amount))
            pcommon = common + ("product", product, "code", code)
            add(product_sales_m, ("period", m) + pcommon, qty=qty, amount=max(0, amount))
            add(product_sales_w, ("period", w) + pcommon, qty=qty, amount=max(0, amount))
            try:
                d = date.fromisocalendar(year, week, 7)
                sales_max_date = max(sales_max_date or d, d)
            except ValueError:
                pass

    consume_sales("历史", legacy=True)
    consume_sales("数据源", legacy=False)

    issues_m, issues_w = defaultdict(dict), defaultdict(dict)
    comp_m, comp_w = defaultdict(dict), defaultdict(dict)
    audit = []
    unmatched = low_conf = duplicate = valid_issues = 0
    after_min = after_max = None

    for sheet_name in ("25年售后登记", "26年售后登记"):
        ws = after_wb[sheet_name]
        headers = [text(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
        for row_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            r = row_dict(headers, values)
            d = parse_date(r.get("⚡️登记日期"))
            if not d:
                continue
            if "重复" in text(r.get("🔐查重标记")):
                duplicate += 1
                continue
            m, w = period_keys(d)
            product_name = text(r.get("产品名称"))
            code = name_to_code.get(product_name, "")
            info = master.get(code)
            unmatched_row = not bool(info)
            if not info:
                info = {
                    "code": code,
                    "cat1": text(r.get("⚡️一级分类")) or "未分类",
                    "cat2": text(r.get("⚡️二级分类")) or "未分类",
                    "product": product_name or "未匹配产品",
                    "barcode": "",
                    "spec": "",
                }
            issue, confidence = classify(r)
            common_m = ("period", m, "cat1", info["cat1"], "cat2", info["cat2"], "issue", issue,
                        "product", info["product"], "code", info.get("code", code), "confidence", confidence)
            common_w = ("period", w, "cat1", info["cat1"], "cat2", info["cat2"], "issue", issue,
                        "product", info["product"], "code", info.get("code", code), "confidence", confidence)
            record_type = text(r.get("🗂登记类型"))
            amount = number(r.get("💰补偿 / 打款金额"))
            store = text(r.get("店铺")) or "未填写"
            platform = aftersales_platform(store)
            skey = shop_key(store, platform)
            paid = bool(parse_date(r.get("✅打款时间")))
            payment_status = "已打款" if paid else "未打款"
            if record_type == "售后登记":
                if unmatched_row:
                    unmatched += 1
                if confidence == "低":
                    low_conf += 1
                add(issues_m, common_m, count=1, amount=amount if 0 < amount <= 100000 else 0)
                add(issues_w, common_w, count=1, amount=amount if 0 < amount <= 100000 else 0)
                valid_issues += 1
                after_min = min(after_min or d, d)
                after_max = max(after_max or d, d)
            if amount > 0:
                if amount > 100000:
                    audit.append({
                        "date": d.isoformat(), "periodMonth": m, "periodWeek": w,
                        "platform": platform, "shop": store, "shopKey": skey,
                        "recordType": record_type or "未填写", "paymentStatus": payment_status,
                        "cat1": info["cat1"], "cat2": info["cat2"],
                        "issue": issue, "product": info["product"], "code": info.get("code", code),
                        "amount": amount, "reason": "单笔金额超过10万元，疑似小数点/录入错误",
                        "sourceRow": f"{sheet_name}!{row_no}",
                    })
                else:
                    detail = ("platform", platform, "shop", store, "shopKey", skey,
                              "recordType", record_type or "未填写", "paymentStatus", payment_status)
                    cm = common_m + detail
                    cw = common_w + detail
                    add(comp_m, cm, count=1, amount=amount, paid_amount=amount if paid else 0,
                        paid_count=1 if paid else 0, max_single=amount)
                    add(comp_w, cw, count=1, amount=amount, paid_amount=amount if paid else 0,
                        paid_count=1 if paid else 0, max_single=amount)

    data = {
        "meta": {
            "generatedAt": datetime.now().isoformat(timespec="seconds"),
            "afterSalesMin": after_min.isoformat() if after_min else "",
            "afterSalesMax": after_max.isoformat() if after_max else "",
            "salesMax": sales_max_date.isoformat() if sales_max_date else "",
            "validIssues": valid_issues,
            "duplicatesExcluded": duplicate,
            "unmatchedProducts": unmatched,
            "lowConfidence": low_conf,
            "classificationVersion": "售后原因规则-v1.1-20260821",
            "compensationLimit": 100000,
            "salesScope": "电商渠道（已排除经销商）",
            "excludedDealerQty": round(excluded_dealer_qty, 4),
            "excludedDealerAmount": round(excluded_dealer_amount, 6),
        },
        "salesMonth": compact(sales_m), "salesWeek": compact(sales_w),
        "issuesMonth": compact(issues_m), "issuesWeek": compact(issues_w),
        "compMonth": compact(comp_m), "compWeek": compact(comp_w),
        "compensationAudit": audit,
    }
    product_data = {"productSalesMonth": compact(product_sales_m), "productSalesWeek": compact(product_sales_w)}
    shop_sales_data = {"shopSalesMonth": compact(shop_sales_m), "shopSalesWeek": compact(shop_sales_w)}
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text("window.AFTER_SALES_DATA=" + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    PRODUCT_OUTPUT.write_text("window.AFTER_SALES_PRODUCT_SALES=" + json.dumps(product_data, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    SHOP_SALES_OUTPUT.write_text("window.AFTER_SALES_SHOP_SALES=" + json.dumps(shop_sales_data, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    print(json.dumps({
        "output": str(OUTPUT), "bytes": OUTPUT.stat().st_size,
        "issues": valid_issues, "duplicates": duplicate, "unmatched": unmatched,
        "lowConfidence": low_conf, "audit": len(audit),
        "salesMonthRows": len(data["salesMonth"]), "productSalesMonthRows": len(product_data["productSalesMonth"]),
        "compTotal": round(sum(r.get("amount", 0) for r in data["compMonth"]), 2),
        "compPaidTotal": round(sum(r.get("paidAmount", 0) for r in data["compMonth"]), 2),
        "excludedDealerQty": round(excluded_dealer_qty, 2),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
